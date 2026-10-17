import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    # 16:9 widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6] # completely blank layout

    # Color Palette (Preserving visual identity: Cream/Light background, Navy blue headers, Orange accents)
    BG_COLOR = RGBColor(248, 249, 250)         # Soft off-white / light slate
    NAVY_PRIMARY = RGBColor(16, 44, 87)        # #102C57 Deep Navy
    ROYAL_BLUE = RGBColor(30, 86, 160)         # #1E56A0 Royal Blue
    ORANGE_ACCENT = RGBColor(235, 130, 60)     # #EB823C Orange Accent
    RED_ACCENT = RGBColor(219, 68, 85)         # #DB4455 Red Accent
    CARD_BG = RGBColor(255, 255, 255)          # White card
    CARD_BORDER = RGBColor(220, 226, 235)      # Soft gray border
    TEXT_DARK = RGBColor(30, 41, 59)           # Slate 800
    TEXT_MUTED = RGBColor(100, 116, 139)       # Slate 500
    GREEN_ACCENT = RGBColor(34, 153, 84)       # Emerald Green

    # Helper function to add a card shape
    def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1.5)
        return shape

    # Helper function to add header bar
    def add_top_bar(slide, title_text, team_text="Coder Janta Party (CJP)", page_num=2):
        # Team pill on top-left
        pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(0.4), Inches(2.2), Inches(0.6))
        pill.fill.solid()
        pill.fill.fore_color.rgb = RGBColor(240, 244, 250)
        pill.line.color.rgb = ROYAL_BLUE
        tf = pill.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = team_text
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = NAVY_PRIMARY
        p.alignment = PP_ALIGN.CENTER

        # Center Title
        txBox = slide.shapes.add_textbox(Inches(2.9), Inches(0.35), Inches(8.5), Inches(0.7))
        tf = txBox.text_frame
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(20)
        p.font.bold = True
        p.font.color.rgb = NAVY_PRIMARY
        p.alignment = PP_ALIGN.CENTER

        # Footer
        footerBox = slide.shapes.add_textbox(Inches(0.5), Inches(7.0), Inches(12.333), Inches(0.4))
        tf = footerBox.text_frame
        p = tf.paragraphs[0]
        p.text = f"@SIH Idea submission  •  Page {page_num}"
        p.font.size = Pt(10)
        p.font.color.rgb = TEXT_MUTED
        p.alignment = PP_ALIGN.CENTER

    # =========================================================================
    # SLIDE 1: TITLE SLIDE
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)

    # Big Top Header
    titleBox = s1.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.5), Inches(0.8))
    tf = titleBox.text_frame
    p = tf.paragraphs[0]
    p.text = "SMART INDIA HACKATHON 2026"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = NAVY_PRIMARY
    p.alignment = PP_ALIGN.CENTER

    # Project Subtitle
    subBox = s1.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(11.5), Inches(0.8))
    tf = subBox.text_frame
    p = tf.paragraphs[0]
    p.text = "TraceLeak — Cryptographic Attribution for Decrypted Documents"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = ROYAL_BLUE
    p.alignment = PP_ALIGN.CENTER

    # Left Metadata Card
    add_card(s1, Inches(0.8), Inches(2.2), Inches(7.5), Inches(4.5))
    metaBox = s1.shapes.add_textbox(Inches(1.1), Inches(2.35), Inches(7.0), Inches(4.2))
    tf = metaBox.text_frame
    tf.word_wrap = True

    items = [
        ("Problem Statement ID:", "SIH26237"),
        ("Problem Statement Title:", "Cryptographic Attribution and Immutable Decryption Provenance for Multi-Recipient Encrypted Document Distribution"),
        ("Organization / Ministry:", "Ministry of Defence — Weapons and Electronics Systems Engineering Establishment (WESEE)"),
        ("Theme:", "Blockchain & Cybersecurity"),
        ("PS Category:", "Software"),
        ("Team ID:", "121538"),
        ("Team Name:", "Coder Janta Party (CJP)")
    ]

    for i, (label, val) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(8)
        run1 = p.add_run()
        run1.text = f"•  {label} "
        run1.font.bold = True
        run1.font.size = Pt(13)
        run1.font.color.rgb = NAVY_PRIMARY

        run2 = p.add_run()
        run2.text = val
        run2.font.bold = (label in ["Organization / Ministry:", "Team Name:"])
        run2.font.size = Pt(13)
        run2.font.color.rgb = ROYAL_BLUE if label == "Organization / Ministry:" else TEXT_DARK

    # Right Accent Card (Brain / Security Badge)
    rightCard = add_card(s1, Inches(8.7), Inches(2.2), Inches(3.8), Inches(4.5), bg_color=RGBColor(240, 244, 250), border_color=ROYAL_BLUE)
    rcBox = s1.shapes.add_textbox(Inches(8.9), Inches(2.5), Inches(3.4), Inches(3.9))
    tf = rcBox.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "DEFENCE-GRADE SECURITY"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = NAVY_PRIMARY
    p.alignment = PP_ALIGN.CENTER
    p.space_after = Pt(14)

    badges = [
        "🛡️ NIST Post-Quantum Standards\n   FIPS 203 (ML-KEM) & FIPS 204 (ML-DSA)",
        "🔒 100% Air-Gapped Operation\n   Zero reliance on public cloud or internet",
        "⚖️ Indian Evidence Act Ready\n   Sec. 65B & BSA Sec. 63 Admissibility",
        "⚡ Sub-Second Verification\n   < 430 ms end-to-end attribution"
    ]
    for b in badges:
        p = tf.add_paragraph()
        p.text = b
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(12)

    # =========================================================================
    # SLIDE 2: REAL-WORLD ISSUE & METHODOLOGY
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    add_top_bar(s2, "TraceLeak: Cryptographic Attribution for Decrypted Documents", page_num=2)

    # 1. REAL-WORLD ISSUE CARD (Top Left)
    add_card(s2, Inches(0.5), Inches(1.1), Inches(4.3), Inches(1.8))
    # Badge
    b1 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.7), Inches(1.2), Inches(1.8), Inches(0.35))
    b1.fill.solid()
    b1.fill.fore_color.rgb = RED_ACCENT
    b1.line.color.rgb = RED_ACCENT
    b1.text_frame.paragraphs[0].text = "REAL-WORLD ISSUE"
    b1.text_frame.paragraphs[0].font.size = Pt(10)
    b1.text_frame.paragraphs[0].font.bold = True
    b1.text_frame.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
    b1.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

    t1 = s2.shapes.add_textbox(Inches(0.7), Inches(1.6), Inches(3.9), Inches(1.2))
    tf1 = t1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "Command encrypts one classified file for multiple officers. When decrypted, all copies are visually identical. If leaked, every recipient has plausible deniability, and mutable server logs cannot prove culpability in court."
    p.font.size = Pt(10.5)
    p.font.color.rgb = TEXT_DARK

    # 2. WHAT IT DOES CARD (Bottom Left)
    add_card(s2, Inches(0.5), Inches(3.05), Inches(4.3), Inches(3.8))
    b2 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.7), Inches(3.15), Inches(1.6), Inches(0.35))
    b2.fill.solid()
    b2.fill.fore_color.rgb = ORANGE_ACCENT
    b2.line.color.rgb = ORANGE_ACCENT
    b2.text_frame.paragraphs[0].text = "WHAT IT DOES?"
    b2.text_frame.paragraphs[0].font.size = Pt(10)
    b2.text_frame.paragraphs[0].font.bold = True
    b2.text_frame.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
    b2.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

    t2 = s2.shapes.add_textbox(Inches(0.7), Inches(3.55), Inches(3.9), Inches(3.2))
    tf2 = t2.text_frame
    tf2.word_wrap = True

    wid_items = [
        ("1. Dynamic Watermarking", "Invisibly injects frequency-domain 2D Block-DCT QIM mark unique to recipient and session at decryption."),
        ("2. Post-Quantum Signing", "Recipient's local ML-DSA-65 key cryptographically signs the event, enforcing undeniable non-repudiation."),
        ("3. Tamper-Evident Ledger", "Committed to an append-only 4-node notary quorum ledger. No single IT admin can delete or alter history."),
        ("4. Blind Leak Attribution", "Recovers mark from leaked PDF with zero passwords or keys, proving leaker identity deterministically.")
    ]
    for i, (title, desc) in enumerate(wid_items):
        p = tf2.paragraphs[0] if i == 0 else tf2.add_paragraph()
        p.space_after = Pt(6)
        r1 = p.add_run()
        r1.text = title + "\n"
        r1.font.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = NAVY_PRIMARY
        r2 = p.add_run()
        r2.text = desc
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = TEXT_DARK

    # 3. METHODOLOGY & PROCESS (Right Side Flow: 6 steps)
    add_card(s2, Inches(5.0), Inches(1.1), Inches(7.8), Inches(3.7))
    bm = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.2), Inches(1.2), Inches(2.6), Inches(0.35))
    bm.fill.solid()
    bm.fill.fore_color.rgb = ORANGE_ACCENT
    bm.line.color.rgb = ORANGE_ACCENT
    bm.text_frame.paragraphs[0].text = "METHODOLOGY & PROCESS"
    bm.text_frame.paragraphs[0].font.size = Pt(10)
    bm.text_frame.paragraphs[0].font.bold = True
    bm.text_frame.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
    bm.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

    flow_steps = [
        ("01 PROTECT", "Encrypt Once (O(1))", "AES-256 payload encryption; wrap key slots via Post-Quantum ML-KEM-768.", ROYAL_BLUE),
        ("02 OPEN", "Unlock Enclave", "Argon2id vault unlocks locally; generates fresh session nonce and watermark ID.", ROYAL_BLUE),
        ("03 MARK", "Dynamic Watermark", "Injects invisible 2D Block-DCT QIM steganography; copies look visually identical.", GREEN_ACCENT),
        ("04 ATTEST", "Sign Provenance", "NIST ML-DSA-65 private key signs canonical decryption event for non-repudiation.", GREEN_ACCENT),
        ("05 RECORD", "Seal Ledger", "4-node notary quorum verifies and commits receipt with Merkle tree roots.", ORANGE_ACCENT),
        ("06 TRACE", "Blind Forensics", "Extracts watermark from suspect PDF, checks Merkle inclusion, issues evidence dossier.", RED_ACCENT)
    ]

    for idx, (code, step_title, step_desc, col) in enumerate(flow_steps):
        row = idx // 3
        col_idx = idx % 3
        box_left = Inches(5.2 + col_idx * 2.5)
        box_top = Inches(1.7 + row * 1.5)

        card = add_card(s2, box_left, box_top, Inches(2.35), Inches(1.35), bg_color=RGBColor(255, 255, 255), border_color=col)
        tb = s2.shapes.add_textbox(box_left + Inches(0.08), box_top + Inches(0.05), Inches(2.19), Inches(1.25))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = f"{code}: {step_title}"
        p.font.bold = True
        p.font.size = Pt(10)
        p.font.color.rgb = col
        p.space_after = Pt(3)

        p2 = tf.add_paragraph()
        p2.text = step_desc
        p2.font.size = Pt(8.5)
        p2.font.color.rgb = TEXT_DARK

    # 4. INNOVATION CARDS (Bottom Right: 4 Pillars)
    innov_banner = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.0), Inches(4.95), Inches(7.8), Inches(0.35))
    innov_banner.fill.solid()
    innov_banner.fill.fore_color.rgb = ORANGE_ACCENT
    innov_banner.line.color.rgb = ORANGE_ACCENT
    innov_banner.text_frame.paragraphs[0].text = "CORE INNOVATION PILLARS"
    innov_banner.text_frame.paragraphs[0].font.size = Pt(10)
    innov_banner.text_frame.paragraphs[0].font.bold = True
    innov_banner.text_frame.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
    innov_banner.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

    innov_cards = [
        ("FINGERPRINT", "Identifies exact leaked copy\n(2D Block-DCT QIM)", GREEN_ACCENT),
        ("SIGNATURE", "Legally binds recipient identity\n(NIST ML-DSA-65)", ROYAL_BLUE),
        ("OFFLINE LEDGER", "Guarantees tamper-proof history\n(4-Node Quorum)", ORANGE_ACCENT),
        ("AIR-GAPPED", "Runs 100% offline with zero cloud\nor public chain dependencies", RED_ACCENT)
    ]

    for idx, (title, text, col) in enumerate(innov_cards):
        box_left = Inches(5.0 + idx * 1.98)
        card = add_card(s2, box_left, Inches(5.4), Inches(1.88), Inches(1.45), border_color=col)
        tb = s2.shapes.add_textbox(box_left + Inches(0.08), Inches(5.48), Inches(1.72), Inches(1.25))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = title
        p.font.bold = True
        p.font.size = Pt(10)
        p.font.color.rgb = col
        p.alignment = PP_ALIGN.CENTER
        p.space_after = Pt(4)

        p2 = tf.add_paragraph()
        p2.text = text
        p2.font.size = Pt(8.5)
        p2.font.color.rgb = TEXT_DARK
        p2.alignment = PP_ALIGN.CENTER

    # =========================================================================
    # SLIDE 3: TECHNICAL APPROACH & STACK
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    add_top_bar(s3, "TECHNICAL APPROACH", page_num=3)

    # Left Tech Stack Card
    add_card(s3, Inches(0.5), Inches(1.1), Inches(6.8), Inches(5.7))
    b3 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.7), Inches(1.2), Inches(2.6), Inches(0.35))
    b3.fill.solid()
    b3.fill.fore_color.rgb = ORANGE_ACCENT
    b3.line.color.rgb = ORANGE_ACCENT
    b3.text_frame.paragraphs[0].text = "TECHNOLOGIES USED"
    b3.text_frame.paragraphs[0].font.size = Pt(10)
    b3.text_frame.paragraphs[0].font.bold = True
    b3.text_frame.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
    b3.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

    t3 = s3.shapes.add_textbox(Inches(0.7), Inches(1.65), Inches(6.4), Inches(5.0))
    tf3 = t3.text_frame
    tf3.word_wrap = True

    tech_specs = [
        ("FRONTEND", "Cross-platform Electron desktop app with React 19 + Tailwind CSS (Role-based Command, Recipient, and Forensics Consoles)."),
        ("BACKEND & RUNTIME", "Python 3.12 with FastAPI local enclave daemon, backed by SQLite & PostgreSQL for 100% offline persistence and multi-station sync."),
        ("POST-QUANTUM CRYPTOGRAPHY", "AES-256-GCM (Payload AEAD), Argon2id (64MB memory-hard vault), NIST FIPS 203 ML-KEM-768 (Quantum key encapsulation), NIST FIPS 204 ML-DSA-65 (Quantum digital signatures)."),
        ("FORENSIC WATERMARKING", "2D Block-DCT QIM frequency-domain embedding via NumPy & SciPy, plus PyMuPDF/pikepdf structural dictionary carriers for screenshot and re-compression resilience."),
        ("PROVENANCE LEDGER", "Append-only permissioned ledger with SHA-256 block-chaining, Merkle tree inclusion proofs, and 4-node notary multi-signature quorum.")
    ]

    for i, (label, val) in enumerate(tech_specs):
        p = tf3.paragraphs[0] if i == 0 else tf3.add_paragraph()
        p.space_after = Pt(10)
        r1 = p.add_run()
        r1.text = label + ": "
        r1.font.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = NAVY_PRIMARY

        r2 = p.add_run()
        r2.text = val
        r2.font.size = Pt(10)
        r2.font.color.rgb = TEXT_DARK

    # Right Architecture Overview Card
    add_card(s3, Inches(7.5), Inches(1.1), Inches(5.3), Inches(5.7), bg_color=RGBColor(240, 244, 250), border_color=ROYAL_BLUE)
    t3r = s3.shapes.add_textbox(Inches(7.7), Inches(1.25), Inches(4.9), Inches(5.4))
    tf3r = t3r.text_frame
    tf3r.word_wrap = True

    p = tf3r.paragraphs[0]
    p.text = "END-TO-END TRUST ARCHITECTURE"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = NAVY_PRIMARY
    p.space_after = Pt(10)

    arch_points = [
        ("🔐 In-Enclave Decryption (127.0.0.1)", "All passwords, private keys, and plaintexts remain strictly inside the local recipient machine enclave. Zero plaintexts or keys over external networks."),
        ("⚡ O(1) Broadcast Storage", "Command stores and broadcasts exactly 1 single encrypted file to all N recipients. Dynamic watermark is synthesized in-enclave without server bloat."),
        ("📜 Proof of Non-Tampering", "Each receipt is sealed with a 4-node notary quorum. Even if a local workstation database is tampered with, the cryptographic Merkle root flags the discrepancy."),
        ("⚖️ Court-Admissible Dossier", "Clicking 'Export Evidence' generates an automated PDF dossier with digital signatures and inclusion proofs under Sec. 65B Indian Evidence Act / BSA Sec. 63.")
    ]

    for title, desc in arch_points:
        p = tf3r.add_paragraph()
        p.space_after = Pt(8)
        r1 = p.add_run()
        r1.text = title + "\n"
        r1.font.bold = True
        r1.font.size = Pt(10.5)
        r1.font.color.rgb = ROYAL_BLUE

        r2 = p.add_run()
        r2.text = desc
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = TEXT_DARK

    # Output file
    desktop_path = os.path.expanduser(r"~\Desktop\TraceLeak_SIH26237_Updated_Presentation.pptx")
    prs.save(desktop_path)
    print(f"Presentation created successfully at: {desktop_path}")

if __name__ == '__main__':
    create_deck()
